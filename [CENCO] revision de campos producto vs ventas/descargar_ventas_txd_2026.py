"""Descarga todas las ventas TXD de 2026 para los SKU de productos TXD 2026.

No calcula ni filtra el último mes por SKU. Sólo filtra ventas cuyo mes empieza
por 2026-, por ejemplo 2026-01 a 2026-12.
"""

from __future__ import annotations

from pathlib import Path

import xlsxwriter
from openpyxl import load_workbook

try:
    from tqdm import tqdm
except ImportError:
    class tqdm:
        def __init__(self, total=None, **_kwargs): self.total, self.count = total, 0
        def __enter__(self): return self
        def __exit__(self, *_args): return False
        def update(self, value):
            self.count += value
            print(f"SKU procesados: {self.count}/{self.total}")

from odoo_jsonrpc_client import client_from_environment


INPUT = Path("productos_txd_2026.xlsx")
OUTPUT = Path("ventas_txd_2026.xlsx")
MODEL = "venta_txd_raw"
BATCH_SIZE = 10_000
PAGE_SIZE = 5_000
MAX_DATA_ROWS = 999_999
HEADERS = [
    "sku_unidad_negocio", "id_venta_txd", "nombre_producto_venta", "periodo_venta",
    "codigo_tipo_marca", "tipo_marca", "marca_venta", "origen_venta",
]


def display_value(value) -> str:
    """Convierte campos simples y relaciones many2one a texto para Excel."""
    if isinstance(value, list) and len(value) > 1:
        return str(value[1] or "")
    return "" if value is None else str(value)


def product_skus() -> list[str]:
    workbook = load_workbook(INPUT, read_only=True, data_only=True)
    skus: set[str] = set()
    for sheet in workbook.worksheets:
        header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        column = header.index("sku_unidad_negocio")
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row[column] is not None and str(row[column]).strip():
                skus.add(str(row[column]).strip())
    return sorted(skus)


def chunks(items: list[str], size: int):
    for start in range(0, len(items), size):
        yield items[start:start + size]


def main() -> None:
    if not INPUT.exists():
        raise FileNotFoundError(f"No existe {INPUT}")
    skus = product_skus()
    odoo = client_from_environment()
    odoo.authenticate()

    workbook = xlsxwriter.Workbook(OUTPUT, {"constant_memory": True})
    header_format = workbook.add_format({"bold": True, "font_color": "#FFFFFF", "bg_color": "#1F4E78", "align": "center"})
    text_format = workbook.add_format({"font_name": "Arial", "font_size": 10})
    part, row_number, rows_written = 0, MAX_DATA_ROWS, 0

    with tqdm(total=len(skus), unit="SKU", desc="Descargando ventas TXD 2026") as progress:
        for sku_batch in chunks(skus, BATCH_SIZE):
            offset = 0
            while True:
                sales = odoo.execute_kw(MODEL, "search_read", [[
                    ("codigo_item", "in", sku_batch),
                    ("mes", "=like", "2026-%"),
                ]], {"fields": [
                    "codigo_item", "item", "mes", "codigo_tipo_marca", "tipo_marca",
                    "marca", "origen_producto",
                ], "limit": PAGE_SIZE, "offset": offset, "order": "id"})
                if not sales:
                    break
                offset += len(sales)
                for sale in sales:
                    if row_number >= MAX_DATA_ROWS:
                        part += 1; row_number = 0
                        sheet = workbook.add_worksheet(f"Ventas TXD 2026 {part}")
                        sheet.hide_gridlines(2); sheet.freeze_panes(1, 0)
                        sheet.write_row(0, 0, HEADERS, header_format)
                        sheet.set_column("A:B", 18, text_format); sheet.set_column("C:C", 42, text_format)
                        sheet.set_column("D:H", 20, text_format)
                    sheet.write_row(row_number + 1, 0, [
                        sale["codigo_item"] or "", sale["id"], sale["item"] or "", sale["mes"] or "",
                        display_value(sale.get("codigo_tipo_marca")),
                        display_value(sale.get("tipo_marca")),
                        display_value(sale.get("marca")),
                        display_value(sale.get("origen_producto")),
                    ], text_format)
                    row_number += 1; rows_written += 1
            progress.update(len(sku_batch))
    workbook.close()
    print(f"Archivo creado: {OUTPUT.resolve()} | Ventas: {rows_written} | Hojas: {part}")


if __name__ == "__main__":
    main()
