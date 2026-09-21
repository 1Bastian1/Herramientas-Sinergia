"""Cruza productos MDH con la venta MDH Raw más reciente por artículo."""
from __future__ import annotations

import re
import unicodedata
from pathlib import Path
import xlsxwriter
from openpyxl import load_workbook

PRODUCTS = Path(__file__).resolve().parent / "productos_mdh_2026.xlsx"
SALES = Path(__file__).resolve().parent / "ventas_mdh_raw_2026.xlsx"
OUTPUT = Path(__file__).resolve().parent / "comparacion_mdh_2026.xlsx"


def rows(path):
    book = load_workbook(path, read_only=True, data_only=True)
    for sheet in book.worksheets:
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        for values in sheet.iter_rows(min_row=2, values_only=True):
            yield dict(zip(headers, values))


def text_key(value) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(c for c in text if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", text).strip().casefold()


def origin_key(value) -> str:
    key = re.sub(r"[^A-Z0-9]", "", text_key(value).upper())
    return {
        "NAC": "NACIONAL", "NACIONAL": "NACIONAL",
        "IMP": "IMPORTADO", "IMPORT": "IMPORTADO", "IMPORTADO": "IMPORTADO",
    }.get(key, key)


def month_number(value) -> int:
    match = re.search(r"\d+", str(value or ""))
    return int(match.group()) if match else -1


def main() -> None:
    products_by_sku = {}
    for product in rows(PRODUCTS):
        sku = str(product.get("sku_unidad_negocio") or "").strip()
        if sku:
            products_by_sku.setdefault(sku, []).append(product)

    latest_by_article = {}
    for sale in rows(SALES):
        article = str(sale.get("articulo") or "").strip()
        if not article:
            continue
        candidate = (month_number(sale.get("mes")), int(sale.get("id") or 0))
        if article not in latest_by_article or candidate > latest_by_article[article][0]:
            latest_by_article[article] = (candidate, sale)

    workbook = xlsxwriter.Workbook(OUTPUT, {"constant_memory": True})
    header = workbook.add_format({"bold": True, "font_color": "#FFFFFF", "bg_color": "#1F4E78", "align": "center", "text_wrap": True})
    text = workbook.add_format({"font_name": "Arial", "font_size": 10})
    headers = ["articulo", "id_producto", "sku_producto", "valor_producto", "valor_producto_normalizado", "id_venta", "mes_venta", "valor_venta", "valor_venta_normalizado"]
    counts = {"Marca distinta": 0, "Origen distinto": 0}
    sheets = {}
    for title in counts:
        sheet = workbook.add_worksheet(title)
        sheet.hide_gridlines(2); sheet.freeze_panes(1, 0)
        sheet.write_row(0, 0, headers, header)
        sheet.set_column("A:B", 18, text); sheet.set_column("C:C", 24, text); sheet.set_column("D:I", 24, text)
        sheets[title] = sheet

    for article, (_, sale) in latest_by_article.items():
        for product in products_by_sku.get(article, []):
            comparisons = [
                ("Marca distinta", product.get("marca") or "", sale.get("nombre_marca") or "", text_key),
                ("Origen distinto", product.get("origen") or "", sale.get("categoria_valorizacion") or "", origin_key),
            ]
            for title, product_value, sale_value, normalize in comparisons:
                if normalize(product_value) == normalize(sale_value):
                    continue
                row = counts[title] + 1
                sheets[title].write_row(row, 0, [
                    article, product.get("id_producto"), product.get("sku_producto") or "",
                    product_value, normalize(product_value), sale.get("id"), sale.get("mes") or "",
                    sale_value, normalize(sale_value),
                ], text)
                counts[title] += 1
    workbook.close()
    print(f"Archivo: {OUTPUT.resolve()} | Marca: {counts['Marca distinta']} | Origen: {counts['Origen distinto']} | Artículos con venta: {len(latest_by_article)}", flush=True)


if __name__ == "__main__":
    main()
