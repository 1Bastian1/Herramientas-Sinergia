"""Actualiza el many2one ``producto_maestro`` desde el Excel de revisión.

Por seguridad, el modo inicial procesa solo los dos primeros registros válidos
y requiere ``--confirmar`` antes de escribir en Odoo.
"""

from __future__ import annotations

import argparse
from collections import defaultdict
from dataclasses import dataclass
from getpass import getpass
from pathlib import Path
from time import sleep

from openpyxl import load_workbook

from odoo_jsonrpc import OdooJsonRpcClient, OdooJsonRpcError


MODEL = "levantamiento_rep.producto"
MASTER_FIELD = "producto_maestro"
DEFAULT_FILE = Path(r"C:\Users\Kyokay_lml\Downloads\Revisión de PRODUCTO MAESTRO faltante.xlsx")
SOURCE_ID_COLUMN = "ID"
MASTER_ID_COLUMN = "ID PRODUCTO MAESTRO"
DEFAULT_LIMIT = 2


@dataclass(frozen=True)
class Update:
    product_id: int
    master_id: int
    sku: str
    master_sku: str
    excel_row: int


def required_text(prompt: str, *, password: bool = False) -> str:
    value = getpass(prompt) if password else input(prompt).strip()
    if not value:
        raise ValueError(f"Falta un dato obligatorio: {prompt.rstrip(': ')}")
    return value


def integer(value: object, column: str, row: int) -> int:
    try:
        return int(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Fila {row}: {column} debe tener un ID numérico; se recibió {value!r}.") from exc


def read_updates(path: Path, limit: int | None) -> list[Update]:
    if not path.exists():
        raise FileNotFoundError(f"No se encuentra el Excel: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    required_columns = {SOURCE_ID_COLUMN, MASTER_ID_COLUMN, "SKU Producto", "Producto maestro"}
    missing = required_columns.difference(headers)
    if missing:
        raise ValueError(f"Faltan columnas requeridas en el Excel: {', '.join(sorted(missing))}.")
    positions = {header: index for index, header in enumerate(headers)}

    updates: list[Update] = []
    seen_sources: dict[int, int] = {}
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        source_value = row[positions[SOURCE_ID_COLUMN]]
        master_value = row[positions[MASTER_ID_COLUMN]]
        if source_value in (None, "") and master_value in (None, ""):
            continue
        source_id = integer(source_value, SOURCE_ID_COLUMN, row_number)
        master_id = integer(master_value, MASTER_ID_COLUMN, row_number)
        if source_id == master_id:
            raise ValueError(f"Fila {row_number}: el producto y su maestro no pueden ser el mismo ID ({source_id}).")
        if source_id in seen_sources:
            raise ValueError(f"Fila {row_number}: el producto ID {source_id} ya aparece en la fila {seen_sources[source_id]}.")
        seen_sources[source_id] = row_number
        updates.append(Update(
            product_id=source_id,
            master_id=master_id,
            sku=str(row[positions["SKU Producto"]] or ""),
            master_sku=str(row[positions["Producto maestro"]] or ""),
            excel_row=row_number,
        ))
        if limit is not None and len(updates) >= limit:
            break
    if not updates:
        raise ValueError("El Excel no contiene filas válidas para actualizar.")
    return updates


def validate_records(client: OdooJsonRpcClient, updates: list[Update]) -> None:
    product_ids = [update.product_id for update in updates]
    master_ids = [update.master_id for update in updates]
    existing_products = client.execute_kw(MODEL, "search", [[("id", "in", product_ids)]])
    existing_masters = client.execute_kw(MODEL, "search", [[("id", "in", master_ids)]])
    missing_products = sorted(set(product_ids).difference(existing_products))
    missing_masters = sorted(set(master_ids).difference(existing_masters))
    if missing_products or missing_masters:
        details = []
        if missing_products:
            details.append(f"productos no encontrados: {missing_products}")
        if missing_masters:
            details.append(f"maestros no encontrados: {missing_masters}")
        raise ValueError("; ".join(details))


def show_plan(updates: list[Update]) -> None:
    print(f"Modelo: {MODEL}; campo: {MASTER_FIELD}")
    print(f"Registros preparados: {len(updates)}")
    preview = updates if len(updates) <= 10 else updates[:10]
    for update in preview:
        print(
            f"  Fila {update.excel_row}: producto {update.product_id} ({update.sku}) "
            f"-> maestro {update.master_id} ({update.master_sku})"
        )
    if len(preview) < len(updates):
        print(f"  ... y {len(updates) - len(preview)} asignaciones adicionales.")


def chunks(values: list[int], size: int = 500):
    for start in range(0, len(values), size):
        yield values[start:start + size]


def apply_updates(client: OdooJsonRpcClient, updates: list[Update]) -> None:
    """Escribe grupos con el mismo maestro. Reintentar es seguro e idempotente."""
    grouped: dict[int, list[int]] = defaultdict(list)
    for update in updates:
        grouped[update.master_id].append(update.product_id)

    batches = [
        (master_id, product_ids)
        for master_id, ids in grouped.items()
        for product_ids in chunks(ids)
    ]
    completed = 0
    for batch_number, (master_id, product_ids) in enumerate(batches, start=1):
        for attempt in range(1, 4):
            try:
                client.execute_kw(MODEL, "write", [product_ids, {MASTER_FIELD: master_id}])
                break
            except OdooJsonRpcError:
                if attempt == 3:
                    raise
                sleep(attempt * 2)
        completed += len(product_ids)
        if batch_number == 1 or batch_number % 10 == 0 or batch_number == len(batches):
            print(f"Progreso: {completed}/{len(updates)} productos actualizados ({batch_number}/{len(batches)} lotes).")


def verify_updates(client: OdooJsonRpcClient, updates: list[Update]) -> None:
    expected = {update.product_id: update.master_id for update in updates}
    actual: dict[int, int | bool] = {}
    for product_ids in chunks(list(expected)):
        records = client.execute_kw(MODEL, "read", [product_ids], {"fields": [MASTER_FIELD]})
        for record in records:
            value = record.get(MASTER_FIELD)
            actual[record["id"]] = value[0] if isinstance(value, list) else value
    incorrect = [product_id for product_id, master_id in expected.items() if actual.get(product_id) != master_id]
    if incorrect:
        raise ValueError(f"La verificación falló para {len(incorrect)} producto(s). Primeros IDs: {incorrect[:10]}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Actualiza producto_maestro desde un Excel.")
    parser.add_argument("--archivo", type=Path, default=DEFAULT_FILE, help="Ruta al archivo Excel de revisión.")
    parser.add_argument("--limite", type=int, default=DEFAULT_LIMIT, help="Número de filas a procesar. Por defecto: 2.")
    parser.add_argument("--todos", action="store_true", help="Procesa todas las filas válidas del Excel.")
    parser.add_argument("--confirmar", action="store_true", help="Autoriza la escritura en Odoo.")
    args = parser.parse_args()
    if args.todos and args.limite != DEFAULT_LIMIT:
        parser.error("Usa solo --todos o --limite, no ambos.")
    if args.limite < 1:
        parser.error("--limite debe ser mayor que cero.")

    updates = read_updates(args.archivo, None if args.todos else args.limite)
    show_plan(updates)
    if not args.confirmar:
        print("Vista previa terminada. No se escribió ningún cambio. Añade --confirmar para aplicar el lote mostrado.")
        return

    print("Ingresa las credenciales de Odoo. La contraseña no se mostrará.")
    client = OdooJsonRpcClient(
        required_text("URL de Odoo: "),
        required_text("Base de datos: "),
        required_text("Usuario: "),
        required_text("Contraseña: ", password=True),
    )
    client.authenticate()
    validate_records(client, updates)
    apply_updates(client, updates)
    verify_updates(client, updates)
    print(f"Proceso terminado: {len(updates)} producto(s) actualizado(s).")


if __name__ == "__main__":
    try:
        main()
    except (OdooJsonRpcError, OSError, ValueError) as error:
        print(f"Proceso detenido: {error}")
        raise SystemExit(1)
